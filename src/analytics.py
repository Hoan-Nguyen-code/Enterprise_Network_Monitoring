
from pathlib import Path
from datetime import datetime

import matplotlib
matplotlib.use("Agg")

import matplotlib.pyplot as plt
import pandas as pd

from openpyxl.styles import (
    Font,
    PatternFill,
    Alignment
)
from openpyxl.chart import LineChart, Reference
from openpyxl.utils import get_column_letter


BASE_DIR = Path(__file__).resolve().parent.parent

DATA_DIR = BASE_DIR / "data"
LOG_DIR = BASE_DIR / "logs"
REPORT_DIR = BASE_DIR / "reports"
CHART_DIR = REPORT_DIR / "charts"

REPORT_DIR.mkdir(parents=True, exist_ok=True)
CHART_DIR.mkdir(parents=True, exist_ok=True)


def load_csv(path):
    if not path.exists() or path.stat().st_size == 0:
        return pd.DataFrame()

    try:
        return pd.read_csv(path)
    except (pd.errors.EmptyDataError, OSError) as exc:
        print(f"Cannot read {path.name}: {exc}")
        return pd.DataFrame()


def prepare_data():
    devices = load_csv(DATA_DIR / "devices.csv")
    interfaces = load_csv(DATA_DIR / "interfaces.csv")
    alerts = load_csv(LOG_DIR / "alerts.csv")

    if not devices.empty:
        devices["timestamp"] = pd.to_datetime(
            devices["timestamp"],
            errors="coerce"
        )

        for column in [
            "cpu_5s",
            "cpu_1m",
            "cpu_5m",
            "ram_percent"
        ]:
            if column in devices.columns:
                devices[column] = pd.to_numeric(
                    devices[column],
                    errors="coerce"
                )

    if not interfaces.empty:
        interfaces["timestamp"] = pd.to_datetime(
            interfaces["timestamp"],
            errors="coerce"
        )

        for column in [
            "rx_bps",
            "tx_bps",
            "input_errors",
            "output_errors",
            "crc_errors"
        ]:
            if column in interfaces.columns:
                interfaces[column] = pd.to_numeric(
                    interfaces[column],
                    errors="coerce"
                )

        interfaces["rx_mbps"] = (
            interfaces["rx_bps"] / 1_000_000
        )

        interfaces["tx_mbps"] = (
            interfaces["tx_bps"] / 1_000_000
        )

    if not alerts.empty:
        alerts["timestamp"] = pd.to_datetime(
            alerts["timestamp"],
            errors="coerce"
        )

    return devices, interfaces, alerts


def create_device_summary(devices):
    if devices.empty:
        return pd.DataFrame()

    summary = devices.groupby(
        "device", as_index=False
    ).agg(
        samples=("device", "size"),
        cpu_average=("cpu_5s", "mean"),
        cpu_max=("cpu_5s", "max"),
        ram_average=("ram_percent", "mean"),
        ram_max=("ram_percent", "max")
    )

    latest = (
        devices.sort_values("timestamp")
        .groupby("device", as_index=False)
        .tail(1)
    )

    latest = latest[
        ["device", "host", "reachable", "timestamp"]
    ].rename(
        columns={
            "reachable": "last_reachable",
            "timestamp": "last_seen"
        }
    )

    summary = summary.merge(
        latest,
        on="device",
        how="left"
    )

    return summary.round(2)


def create_interface_summary(interfaces):
    if interfaces.empty:
        return pd.DataFrame()

    summary = interfaces.groupby(
        ["device", "name"], as_index=False
    ).agg(
        samples=("name", "size"),
        rx_average_mbps=("rx_mbps", "mean"),
        rx_max_mbps=("rx_mbps", "max"),
        tx_average_mbps=("tx_mbps", "mean"),
        tx_max_mbps=("tx_mbps", "max"),
        max_input_errors=("input_errors", "max"),
        max_output_errors=("output_errors", "max"),
        max_crc_errors=("crc_errors", "max")
    )

    return summary.round(4)


def create_alert_summary(alerts):
    if alerts.empty:
        return pd.DataFrame(
            columns=["device", "level", "count"]
        )

    return (
        alerts.groupby(
            ["device", "level"],
            as_index=False
        )
        .size()
        .rename(columns={"size": "count"})
    )


def plot_metric(devices, column, title, filename):
    if devices.empty or column not in devices.columns:
        return

    fig, ax = plt.subplots(figsize=(12, 5))

    plotted = False

    for name, group in devices.groupby("device"):
        group = group.sort_values("timestamp")
        group = group.dropna(
            subset=["timestamp", column]
        )

        if group.empty:
            continue

        ax.plot(
            group["timestamp"],
            group[column],
            marker=".",
            linewidth=1.5,
            label=name
        )

        plotted = True

    if not plotted:
        plt.close(fig)
        return

    ax.set_title(title)
    ax.set_xlabel("Time")
    ax.set_ylabel("Usage (%)")
    ax.set_ylim(0, 100)
    ax.grid(True, alpha=0.3)
    ax.legend()

    fig.autofmt_xdate()
    fig.tight_layout()

    fig.savefig(
        CHART_DIR / filename,
        dpi=160,
        bbox_inches="tight"
    )

    plt.close(fig)


def plot_traffic(interfaces):
    if interfaces.empty:
        return

    fig, ax = plt.subplots(figsize=(12, 5))

    plotted = False

    for (device, name), group in interfaces.groupby(
        ["device", "name"]
    ):
        group = group.sort_values("timestamp")

        for metric, direction in [
            ("rx_mbps", "RX"),
            ("tx_mbps", "TX")
        ]:
            valid = group.dropna(
                subset=["timestamp", metric]
            )

            if valid.empty:
                continue

            ax.plot(
                valid["timestamp"],
                valid[metric],
                linewidth=1.5,
                label=f"{device} {name} {direction}"
            )

            plotted = True

    if not plotted:
        plt.close(fig)
        return

    ax.set_title("Network Traffic History")
    ax.set_xlabel("Time")
    ax.set_ylabel("Traffic (Mbps)")
    ax.grid(True, alpha=0.3)
    ax.legend(fontsize=8)

    fig.autofmt_xdate()
    fig.tight_layout()

    fig.savefig(
        CHART_DIR / "traffic_usage.png",
        dpi=160,
        bbox_inches="tight"
    )

    plt.close(fig)


def prepare_excel_frame(frame):
    result = frame.copy()

    for column in result.columns:
        if pd.api.types.is_datetime64_any_dtype(
            result[column]
        ):
            result[column] = result[column].dt.strftime(
                "%Y-%m-%d %H:%M:%S"
            )

    return result


def style_worksheet(ws):
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    for cell in ws[1]:
        cell.font = Font(
            bold=True,
            color="FFFFFF"
        )

        cell.fill = PatternFill(
            fill_type="solid",
            fgColor="17365D"
        )

        cell.alignment = Alignment(
            horizontal="center"
        )

    for column in ws.columns:
        column_letter = get_column_letter(
            column[0].column
        )

        max_length = max(
            len(str(cell.value))
            if cell.value is not None else 0
            for cell in column
        )

        ws.column_dimensions[
            column_letter
        ].width = min(max(max_length + 3, 12), 38)


def export_excel(
    devices,
    interfaces,
    alerts,
    device_summary,
    interface_summary,
    alert_summary
):
    output = REPORT_DIR / "network_analysis.xlsx"

    sheets = {
        "Device Summary": device_summary,
        "Interface Summary": interface_summary,
        "Alert Summary": alert_summary,
        "Device History": devices,
        "Interface History": interfaces,
        "Alert History": alerts
    }

    with pd.ExcelWriter(
        output,
        engine="openpyxl"
    ) as writer:

        for sheet_name, frame in sheets.items():
            if frame.empty and len(frame.columns) == 0:
                frame = pd.DataFrame({
                    "Information": ["No data available"]
                })

            excel_frame = prepare_excel_frame(frame)

            excel_frame.to_excel(
                writer,
                sheet_name=sheet_name,
                index=False
            )

            ws = writer.sheets[sheet_name]
            style_worksheet(ws)

    return output


def main():
    print("=" * 60)
    print("ENTERPRISE NETWORK DATA ANALYTICS")
    print("=" * 60)

    devices, interfaces, alerts = prepare_data()

    if devices.empty and interfaces.empty:
        print(
            "No monitoring data found. "
            "Run src/main.py first."
        )
        return

    device_summary = create_device_summary(devices)
    interface_summary = create_interface_summary(
        interfaces
    )
    alert_summary = create_alert_summary(alerts)

    plot_metric(
        devices,
        "cpu_5s",
        "CPU Usage History",
        "cpu_usage.png"
    )

    plot_metric(
        devices,
        "ram_percent",
        "RAM Usage History",
        "ram_usage.png"
    )

    plot_traffic(interfaces)

    excel_path = export_excel(
        devices,
        interfaces,
        alerts,
        device_summary,
        interface_summary,
        alert_summary
    )

    print("\nDEVICE SUMMARY")
    if not device_summary.empty:
        print(device_summary.to_string(index=False))

    print("\nINTERFACE SUMMARY")
    if not interface_summary.empty:
        print(interface_summary.to_string(index=False))

    print("\nALERT SUMMARY")
    if not alert_summary.empty:
        print(alert_summary.to_string(index=False))
    else:
        print("No recorded alerts.")

    print("\nReports generated:")
    print(f"Excel: {excel_path}")
    print(f"Charts: {CHART_DIR}")

    print("\nAnalysis completed successfully.")


if __name__ == "__main__":
    main()
